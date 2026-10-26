# -*- coding: utf-8 -*-
"""实体台账导出：从 graph_resolved.json 自动生成全量清单 xlsx。

分表：单位（参与项目/角色汇总/事项责任/人员）· 项目（参与单位/会议文件/时间跨度）·
      人员（所属单位/参会）· 事项（内容/项目/责任单位/期限/出处）。

用法：.venv\\Scripts\\python.exe tools\\export_inventory.py
产物：output\\实体台账_<日期>.xlsx
"""
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from kg.serve.query import build_file_dates, load_graph            # noqa: E402

import openpyxl                                                    # noqa: E402
from openpyxl.styles import Font, PatternFill                      # noqa: E402

GRAPH = ROOT / "var" / "extract" / "graph_resolved.json"
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True)


def main():
    g = load_graph(GRAPH)
    ents, rels = g["entities"], g["relationships"]
    name_of = {e["id"]: e.get("name", "") for e in ents}
    type_of = {e["id"]: e.get("type", "") for e in ents}
    fdates = build_file_dates(ROOT)

    org_proj = defaultdict(set)          # ORG id -> 参与项目名集合
    org_roles = defaultdict(set)         # ORG id -> 角色（项目·日期）
    org_matters = Counter()              # ORG id -> 事项责任数
    org_facts = Counter()                # ORG id -> 关联事实(边)条数
    person_of_org = defaultdict(set)
    person_org = defaultdict(set)
    person_meet = defaultdict(set)
    person_facts = Counter()
    proj_units = defaultdict(set)
    proj_meetings = defaultdict(set)
    proj_files = defaultdict(set)
    matter_proj = {}                     # MATTER id -> 项目名
    matter_unit = defaultdict(list)

    # ⭐ 第一遍：先把 MATTER→PROJECT 的归属建全（避免同轮循环内的顺序依赖）
    for r in rels:
        if r.get("type") == "属于项目" and r.get("source", "").startswith("M::"):
            matter_proj[r["source"]] = name_of.get(r.get("target"), r.get("target"))

    # ⭐ 第二遍：统计（端点按 id 精确匹配）
    for r in rels:
        s, t, ty = r.get("source"), r.get("target"), r.get("type")
        ev = (r.get("properties") or {}).get("evidence") or []
        n_ev = len(ev) if isinstance(ev, list) else 1
        ts, tt = type_of.get(s, ""), type_of.get(t, "")

        if ts == "ORG" or tt == "ORG":
            oid = s if ts == "ORG" else t
            org_facts[oid] += n_ev
        if ts == "PERSON" or tt == "PERSON":
            pid = s if ts == "PERSON" else t
            person_facts[pid] += n_ev

        if ty == "属于项目":
            if s.startswith("M::"):
                pass                                    # matter_proj 第一遍已建
            elif ts == "MEETING" or "MEETING" in str(s):
                proj_meetings[t].add(s)
            for x in (ev if isinstance(ev, list) else [ev]):
                proj_files[t].add(str(x.get("file", "")))
        elif ty == "单位角色":
            proj_units[t].add(name_of.get(s, s))
            org_proj[s].add(name_of.get(t, t))
            for x in (ev if isinstance(ev, list) else [ev]):
                d = fdates.get(Path(str(x.get("file", ""))).name)
                if d:
                    org_roles[s].add("%s（%s）" % (name_of.get(t, t), d))
        elif ty == "事项责任单位":
            matter_unit[s].append(name_of.get(t, t))
            org_matters[t] += 1
            # ⛔ 方向：MATTER→ORG，参与项目记到 t（此前记到 s 导致全 0）
            if s in matter_proj:
                org_proj[t].add(matter_proj[s])
        elif ty == "所属单位":
            person_org[s].add(name_of.get(t, t))
            person_of_org[t].add(name_of.get(s, s))
        elif ty in ("参会", "主持人", "记录"):
            person_meet[s].add(t)

    wb = openpyxl.Workbook()

    # ---- 单位表
    ws = wb.active
    ws.title = "单位"
    ws.append(["序号", "单位名称", "参与项目数", "担任角色（项目·日期）",
               "事项责任数", "所属人员数", "关联事实条数"])
    orgs = [e for e in ents if e.get("type") == "ORG"]
    for i, e in enumerate(sorted(orgs, key=lambda x: -(org_facts.get(x["id"], 0)
                                                       + org_matters.get(x["id"], 0))), 1):
        eid = e["id"]
        roles = "；".join(sorted(org_roles.get(eid, []))[:6])
        ws.append([i, e.get("name"), len([p for p in org_proj.get(eid, []) if p]),
                   roles, org_matters.get(eid, 0),
                   len(person_of_org.get(eid, [])), org_facts.get(eid, 0)])
    ws.append([])
    ws.append(["合计", len(orgs), ""])

    # ---- 项目表
    ws2 = wb.create_sheet("项目")
    ws2.append(["序号", "项目名称", "参与单位数", "参与单位", "关联会议文件数", "时间跨度"])
    projs = [e for e in ents if e.get("type") == "PROJECT"]
    for i, e in enumerate(sorted(projs, key=lambda x: -len(proj_units.get(x["id"], []))), 1):
        eid = e["id"]
        files = proj_files.get(eid, set())
        ds = sorted(fdates.get(Path(f).name) for f in files if Path(f).name in fdates)
        span = "%s ~ %s" % (ds[0], ds[-1]) if ds else ""
        ws2.append([i, e.get("name"), len(proj_units.get(eid, [])),
                    "、".join(sorted(proj_units.get(eid, []))[:8]),
                    len(files), span])

    # ---- 人员表
    ws3 = wb.create_sheet("人员")
    ws3.append(["序号", "姓名", "所属单位", "参会/主持会议数", "关联事实条数"])
    persons = [e for e in ents if e.get("type") == "PERSON"]
    for i, e in enumerate(sorted(persons, key=lambda x: -len(person_meet.get(x["id"], []))), 1):
        eid = e["id"]
        ws3.append([i, e.get("name"), "、".join(sorted(person_org.get(eid, []))[:3]),
                    len(person_meet.get(eid, [])), person_facts.get(eid, 0)])

    # ---- 事项表
    ws4 = wb.create_sheet("事项")
    ws4.append(["序号", "事项内容", "所属项目", "责任单位（原文）", "责任单位（图谱）",
                "期限", "出处文件", "页码", "原文节选"])
    matters = [e for e in ents if e.get("type") == "MATTER"]
    for i, e in enumerate(matters, 1):
        p = e.get("properties") or {}
        ev = p.get("evidence") or []
        if not ev and p.get("file"):
            ev = [{"file": p.get("file"), "page": p.get("page"), "quote": p.get("quote")}]
        x0 = (ev[0] if isinstance(ev, list) else ev) if ev else {}
        ws4.append([i, (e.get("name") or p.get("content") or "")[:80],
                    matter_proj.get(e["id"], ""),
                    p.get("责任单位(原文)", ""), "、".join(matter_unit.get(e["id"], []))[:40],
                    p.get("期限", ""), Path(str(x0.get("file", ""))).name,
                    x0.get("page", ""), (x0.get("quote") or "")[:60]])

    for ws_ in [ws, ws2, ws3, ws4]:
        for c in ws_[1]:
            c.fill, c.font = HDR_FILL, HDR_FONT
        for col, w in zip(ws_.columns, [6, 42, 12, 50, 12, 12, 14, 30, 8, 50]):
            if col:
                ws_.column_dimensions[col[0].column_letter].width = w
        ws_.freeze_panes = "A2"

    out = ROOT / "output" / ("实体台账_%s.xlsx" % date.today().strftime("%Y%m%d"))
    try:
        wb.save(out)
    except PermissionError:
        # 文件被 Excel/预览占用 → 落到带时间的文件名，不打断
        out = ROOT / "output" / ("实体台账_%s_%s.xlsx" % (
            date.today().strftime("%Y%m%d"), __import__("datetime").datetime.now().strftime("%H%M%S")))
        wb.save(out)
    print("台账已生成：", out)
    print("单位 %d | 项目 %d | 人员 %d | 事项 %d" % (len(orgs), len(projs), len(persons), len(matters)))


if __name__ == "__main__":
    main()
